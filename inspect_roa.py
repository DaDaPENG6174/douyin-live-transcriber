from docx import Document
import sys
sys.stdout.reconfigure(encoding='utf-8')
p = r'C:\Users\Administrator\Documents\抖音\转发资料_DOCX\付付Roa-完整逐字稿.docx'
d = Document(p)
for start in [0, 80, 300, 700, 1100, 1400]:
    print(f'\n===== {start} =====')
    for i in range(start, min(start+35, len(d.paragraphs))):
        print(f'[{i}] {d.paragraphs[i].text}')
