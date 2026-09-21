from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(90, 90, 90)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    run.font.name = "等线"
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    run.font.name = "等线"
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "等线"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_text(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True, color=ACCENT)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def build_docx(md_path, out_path):
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "等线"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, ACCENT, 10, 5),
    ]:
        st = styles[name]
        st.font.name = "等线"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "Callout" not in [s.name for s in styles]:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        callout.base_style = normal
    else:
        callout = styles["Callout"]
    callout.font.name = "等线"
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = MUTED
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.18)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = "回避型关系系列直播稿"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run_font(run, size=9, color=MUTED)
    set_page_number(section.footer.paragraphs[0])

    title = next((line[2:].strip() for line in lines if line.startswith("# ")), Path(md_path).stem)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=22, bold=True, color=ACCENT)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(18)
    r = p2.add_run("60 分钟直播口播稿")
    set_run_font(r, size=11, color=MUTED)

    in_code = False
    for line in lines:
        stripped = line.strip()
        if stripped == "```":
            in_code = not in_code
            continue
        if not stripped or stripped == "---" or stripped.startswith("# "):
            continue
        if in_code:
            p = add_text(doc, stripped, style="Callout")
            p.paragraph_format.left_indent = Inches(0.25)
            continue
        m = re.match(r"^(#{2,3})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)) - 1, 3)
            doc.add_paragraph(m.group(2), style=f"Heading {level}")
            continue
        if stripped.startswith(">"):
            add_text(doc, stripped.lstrip("> "), style="Callout")
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(stripped[2:].strip())
            set_run_font(r, size=10.5)
            continue
        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(re.sub(r"^\d+\.\s+", "", stripped))
            set_run_font(r, size=10.5)
            continue
        add_text(doc, stripped)

    core = doc.core_properties
    core.title = title
    core.subject = "回避型关系系列直播稿"
    core.author = ""
    doc.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: md_to_docx.py input.md output.docx")
    build_docx(sys.argv[1], sys.argv[2])
