from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Administrator\Documents\抖音")
OUT = ROOT / "转发资料_DOCX"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C4D6", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_fixed_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def style_document(doc, title):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.18

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 14, 7),
        ("Heading 2", 13, "2E74B5", 11, 5),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "抖音直播话术项目 | 转发资料"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(120, 130, 145)
    footer = section.footer.paragraphs[0]
    footer.text = "内部整理资料"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(140, 140, 140)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title)
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(11, 37, 69)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(55, 65, 81)
    p._p.get_or_add_pPr().append(OxmlElement("w:shd"))
    p._p.pPr[-1].set(qn("w:fill"), "F3F6F9")


def add_markdown(doc, text):
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                add_code(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(line[2:].strip(), style="Heading 1")
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 2")
            i += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:].strip(), style="Heading 3")
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.right_indent = Inches(0.15)
            p.paragraph_format.space_after = Pt(5)
            r = p.add_run(line[2:].strip())
            r.italic = True
            r.font.color.rgb = RGBColor(70, 80, 95)
            i += 1
            continue
        if re.match(r"^[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(re.sub(r"^[-*] ", "", line))
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.add_run(re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|---"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                row = lines[i]
                if set(row.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                    i += 1
                    continue
                cells = [c.strip() for c in row.strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            if rows:
                col_count = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=col_count)
                set_table_borders(table)
                widths = [9360 // col_count] * col_count
                widths[-1] += 9360 - sum(widths)
                set_fixed_table_width(table, widths)
                set_repeat_table_header(table.rows[0])
                for ri, row in enumerate(rows):
                    for ci in range(col_count):
                        cell = table.cell(ri, ci)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                        cell.text = row[ci] if ci < len(row) else ""
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after = Pt(2)
                            for run in p.runs:
                                run.font.size = Pt(8.3 if col_count >= 4 else 9)
                        if ri == 0:
                            set_cell_shading(cell, "E8EEF5")
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.bold = True
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.add_run(line)
        i += 1


def add_transcript(doc, text):
    for raw in text.replace("\r\n", "\n").split("\n"):
        if not raw.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(raw)
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(9.5)


def markdown_doc(src, out):
    doc = Document()
    title = src.stem.replace("-", " ")
    style_document(doc, title)
    add_markdown(doc, src.read_text(encoding="utf-8"))
    doc.save(out)


def transcript_doc(src, out, title):
    doc = Document()
    style_document(doc, title)
    p = doc.add_paragraph("以下为本地 SenseVoiceSmall 转写结果，个别专有名词和数字可能存在识别误差。")
    p.runs[0].italic = True
    p.runs[0].font.color.rgb = RGBColor(100, 110, 125)
    add_transcript(doc, src.read_text(encoding="utf-8"))
    doc.save(out)


def main():
    OUT.mkdir(exist_ok=True)
    for old in OUT.glob("*.docx"):
        old.unlink()
    selected = []
    selected.extend(sorted((ROOT / "docs" / "workflows").glob("*.md")))
    selected.extend(sorted((ROOT / "docs" / "analysis").glob("*.md")))
    selected.extend(sorted((ROOT / "output").glob("unified-analysis.md")))
    selected.extend(sorted(ROOT.glob("output/**/analysis/analysis-report.md")))
    selected.extend(sorted(ROOT.glob("output/**/analysis/fan-acquisition-scripts.md")))
    selected.extend(sorted(ROOT.glob("output/**/analysis/contact-funnel-contexts.md")))

    for src in selected:
        if src.parent.name == "analysis" and src.parents[1].name != "docs":
            source_label = src.parents[1].name
            out_name = f"{source_label}-{src.stem}.docx"
        else:
            out_name = f"{src.stem}.docx"
        out = OUT / out_name
        markdown_doc(src, out)

    transcript_specs = [
        (ROOT / "output/long-recording/2026-08-21_170614/transcript/transcript.txt", "付付Roa-完整逐字稿"),
        (ROOT / "output/recordings/2026-08-21_175230/transcript/transcript.txt", "书书谈回避-完整逐字稿"),
        (ROOT / "output/recordings/2026-08-21_190648/transcript/transcript.txt", "大丽知心说-1小时55分-完整逐字稿"),
        (ROOT / "output/recordings/2026-08-21_162346/transcript/transcript.txt", "大丽知心说-55分钟-完整逐字稿"),
        (ROOT / "output/recordings/2026-08-21_160306/transcript/transcript.txt", "听雨思意-完整逐字稿"),
    ]
    for src, title in transcript_specs:
        transcript_doc(src, OUT / f"{title}.docx", title)

    files = sorted(OUT.glob("*.docx"))
    manifest = OUT / "文件清单.txt"
    manifest.write_text("抖音直播项目 Word 转发资料\n\n" + "\n".join(f"{idx}. {p.name}" for idx, p in enumerate(files, 1)), encoding="utf-8")
    print(f"created {len(files)} docx files in {OUT}")


if __name__ == "__main__":
    main()
