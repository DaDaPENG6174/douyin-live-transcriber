from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"C:\Users\Administrator\Documents\抖音")
BASE = ROOT / "output" / "transcribe" / "2026-08-26_160916"
SRC = BASE / "transcript"
RAW_OUT = ROOT / "2026-08-26_160916_原始逐字稿.docx"
CLEAN_OUT = ROOT / "2026-08-26_160916_整理顺畅版.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def setup(doc, title):
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(14)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(31, 78, 121)
    meta = doc.add_paragraph("来源：2026-08-26_160916.mp4\n识别工具：本地 FunASR / SenseVoiceSmall（CPU）")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in meta.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(100, 100, 100)
    doc.add_paragraph()


def add_body(doc, text, timestamp=None, highlighted=False):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    if timestamp:
        t = p.add_run(f"[{timestamp}] ")
        t.bold = True
        t.font.color.rgb = RGBColor(31, 78, 121)
        t.font.size = Pt(11)
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(14)
    if highlighted:
        r.bold = True


def parse_srt(path):
    raw = path.read_text(encoding="utf-8", errors="replace")
    blocks = re.split(r"\r?\n\r?\n+", raw.strip())
    out = []
    for block in blocks:
        lines = block.splitlines()
        if len(lines) < 3:
            continue
        timing = lines[1]
        text = " ".join(x.strip() for x in lines[2:] if x.strip())
        if text:
            out.append((timing.split(" --> ")[0], text))
    return out


def clean_line(s):
    replacements = {
        "我盒伴都给你": "我陪伴都给你",
        "把命调着就行了": "把命交着就行了",
        "别说我来别说我来粘着你": "别说我来粘着你",
        "我都不你都会觉得": "我都会觉得",
        "我在耽误你": "我在耽误你",
        "自给自足": "自给自足",
        "高能量嘛": "高能量嘛",
        "二8定律": "二八定律",
        "后车": "后撤",
        "充电": "充电",
        "挽留回来": "挽回回来",
        "对你常泽的": "对你负责的",
    }
    for a, b in replacements.items():
        s = s.replace(a, b)
    s = re.sub(r"(啊|嗯|呃|哦)([，。！？、])", r"\2", s)
    s = re.sub(r"([，。！？])\1+", r"\1", s)
    s = re.sub(r"(不是){2,}", "不是", s)
    s = re.sub(r"(对吗){2,}", "对吗", s)
    s = re.sub(r"\s+", "", s)
    return s.strip()


def create_raw():
    doc = Document()
    setup(doc, "2026-08-26 付付Roa 直播录屏｜原始逐字稿")
    note = doc.add_paragraph("说明：以下内容按本地语音识别结果整理，保留原始口语、重复、停顿和识别原貌，并附起始时间。")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(11)
    for stamp, text in parse_srt(SRC / "transcript.srt"):
        add_body(doc, text, stamp)
    doc.save(RAW_OUT)


def create_clean():
    doc = Document()
    setup(doc, "2026-08-26 付付Roa 直播录屏｜整理顺畅版")
    note = doc.add_paragraph("说明：在不改变原文基础意思的前提下，调整明显断句、重复和口头卡顿，保留直播聊天语气，便于阅读和后续口播整理。")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(11)
    for stamp, text in parse_srt(SRC / "transcript.srt"):
        text = clean_line(text)
        if not text or text in {"嗯。", "啊。", "哦。", "呃。"}:
            continue
        # Long ASR chunks are split at sentence punctuation for teleprompter readability.
        pieces = re.split(r"(?<=[。！？])", text)
        pieces = [p.strip() for p in pieces if p.strip()]
        if not pieces:
            continue
        add_body(doc, pieces[0], stamp)
        for piece in pieces[1:]:
            add_body(doc, piece)
    doc.save(CLEAN_OUT)


if __name__ == "__main__":
    create_raw()
    create_clean()
    print(RAW_OUT)
    print(CLEAN_OUT)
